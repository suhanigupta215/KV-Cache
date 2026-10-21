"""Build the editable IDF-B technical annex from measured local evidence."""
from pathlib import Path
import html
import json
import subprocess
from statistics import median

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from reportlab.graphics.shapes import Drawing, String, Rect
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics import renderSVG

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "tmp/disclosure/assets"
ASSETS.mkdir(parents=True, exist_ok=True)
OUT = ROOT / "output/docx/MTR_KV_IDF_B_Detailed.docx"
DATA = ROOT / "results/optimization-20261021"
paired = json.loads((DATA / "paired-serial/manifest.json").read_text())
full = json.loads((DATA / "full-demo-serial-evaluation.json").read_text())
oracle, indexed = (paired["median_end_to_end_ms"][key] for key in ("oracle", "indexed"))
NODE = "/Users/shaurya/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node"
SHARP = "/Users/shaurya/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules/sharp"


def svg_start(width=1200, height=620):
    return [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
            '<rect width="100%" height="100%" fill="white"/>',
            '<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="9" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#26394a"/></marker></defs>']


def text(parts, x, y, lines, size=18, anchor="middle", fill="#172b3a"):
    for index, line in enumerate(lines):
        parts.append(f'<text x="{x}" y="{y + index * (size + 7)}" text-anchor="{anchor}" font-family="Arial" font-size="{size}" fill="{fill}">{html.escape(line)}</text>')


def box(parts, x, y, width, height, lines, fill="#eef4f8", dashed=False):
    parts.append(f'<rect x="{x}" y="{y}" width="{width}" height="{height}" rx="8" fill="{fill}" stroke="#26394a" stroke-width="2" {"stroke-dasharray=\"8 6\"" if dashed else ""}/>')
    text(parts, x + width/2, y + height/2 - (len(lines)-1)*12 + 5, lines,20)


def arrow(parts, points):
    parts.append(f'<polyline points="{points}" fill="none" stroke="#26394a" stroke-width="2" marker-end="url(#arrow)"/>')


def save_svg(name, parts):
    path = ASSETS / (name + ".svg")
    path.write_text("\n".join(parts + ["</svg>"]))
    subprocess.run([NODE, "-e", f"require({json.dumps(SHARP)})({json.dumps(str(path))}).resize(1800).png().toFile({json.dumps(str(path.with_suffix('.png')))})"], check=True)
    return path.with_suffix(".png")


def diagrams():
    s = svg_start(1200, 620)
    nodes = [(30,55,["Encoded video", "H.264 / decoder"]), (270,55,["Normalize", "motion + validity"]),
             (510,55,["MTRE mapping", "crop / patch / page"]), (750,55,["Temporal state", "and scene guard"]),
             (990,55,["MKAD", "generation + tier"]), (990,270,["KMC transaction", "reserve to ready"]),
             (750,270,["Tensor tier store", "T0 / T1 / T2"]), (510,270,["ViT attention", "read stored K/V"]),
             (270,270,["Semantic score", "page attention proxy"]), (30,270,["SRCE correction", "same generation"])]
    for x,y,lines in nodes:
        box(s,x,y,180,95,lines)
    for x in (210,450,690,930): arrow(s,f"{x},102 {x+60},102")
    arrow(s,"1080,150 1080,270")
    for x in (990,750,510,270): arrow(s,f"{x},318 {x-60},318")
    arrow(s,"120,365 120,450 1080,450 1080,365")
    text(s,600,485,["Corrections return through the same generation-checked KMC path"],17)
    box(s,235,525,730,65,["GPU / RTL KMC and full VLM integration remain proposed"],"#ffffff",True)
    save_svg("architecture",s)

    s=svg_start(1200,610)
    s.append('<rect x="40" y="95" width="640" height="360" fill="#f5f5f5" stroke="#26394a" stroke-width="2"/>')
    text(s,360,65,["Source frame 640 x 360 pixels"],23)
    crop=(116*640/455,22.5,224*640/455,315)
    s.append(f'<rect x="{40+crop[0]}" y="{95+crop[1]}" width="{crop[2]}" height="{crop[3]}" fill="#e4eef5" stroke="#156389" stroke-width="3"/>')
    s.append('<rect x="300" y="230" width="85" height="60" fill="#ffe2dc" stroke="#a14231" stroke-width="2"/>')
    text(s,343,258,["codec", "block"],16)
    arrow(s,"520,275 780,275")
    text(s,650,240,["inverse", "preprocessing"],17)
    for row in range(14):
        for col in range(14):
            fill=["#dfebf5","#e7e1f1","#e3edde"][(row*14+col)//32%3]
            s.append(f'<rect x="{820+col*20}" y="{135+row*20}" width="20" height="20" fill="{fill}" stroke="#8f9ca7" stroke-width=".6"/>')
    text(s,960,95,["224 x 224 model input"],22)
    text(s,960,450,["14 x 14 patch grid", "32 patch tokens per page", "7 logical pages"],18)
    text(s,40,505,["Resize 640 x 360 to 455 x 256; center crop starts at (116, 16).", "Source crop: x=163.16, y=22.50, width=315.08, height=315.00.", "Logical pages follow token sequence order and can cross image rows."],19,"start")
    save_svg("geometry",s)

    s=svg_start(1200,560)
    for x,label in [(40,["Validate", "generation + range"]),(330,["Reserve", "replacement blocks"]),(620,["Prepare", "shadow mapping"]),(910,["Commit and ready", "publish then permit read"])]:
        box(s,x,100,250,100,label)
    for x in (290,580,870): arrow(s,f"{x},150 {x+40},150")
    box(s,90,310,430,100,["Stale or duplicate descriptor", "reject without replacing current mapping"],"#faeeee")
    box(s,680,310,430,100,["Replacement reservation fails", "preserve active mapping; fallback / abort"],"#faeeee")
    arrow(s,"165,200 165,310")
    arrow(s,"455,200 455,260 895,260 895,310")
    text(s,600,480,["Generation is a consistency version, not a wall-clock timestamp.", "The Python fence models ordering; no physical GPU memory fence is claimed."],20)
    save_svg("transaction",s)

    s=svg_start(1200,800)
    checks=["Current page mapping committed and ready", "Reliable codec + low activity + STATIC state", "Same controller and consecutive generation", "Same geometry, device, dtype and model parameters", "Payload age within window and exact page input equality"]
    for index,label in enumerate(checks):
        y=35+index*118
        box(s,80,y,690,76,[label],"#eef4f8")
        if index<4:
            arrow(s,f"425,{y+76} 425,{y+118}")
            text(s,450,y+103,["yes"],16)
        if index == 0:
            arrow(s,f"770,{y+38} 840,{y+38}")
            box(s,840,35,330,76,["Wait or fallback", "no model access"],"#faeeee")
        else:
            arrow(s,f"770,{y+38} 880,{y+38} 880,635")
        text(s,805,y+25,["no"],16)
    box(s,690,635,440,90,["Refresh page K/V", "project, encode and store"],"#fff0e5")
    box(s,80,635,490,90,["Reuse payload after generation rebinding", "skip that page's K/V projection"],"#e5f2e7")
    arrow(s,"425,583 425,635")
    text(s,600,775,["Q and prefix K/V are always computed; attention reads all patch K/V through the store."],18)
    save_svg("reuse",s)

    drawing=Drawing(720,350)
    drawing.add(Rect(0,0,720,350,fillColor=__import__('reportlab.lib.colors',fromlist=['white']).white,strokeColor=None))
    chart=VerticalBarChart()
    chart.x,chart.y,chart.width,chart.height=75,65,575,230
    chart.data=[(oracle,indexed)]
    chart.categoryAxis.categoryNames=["Reference oracle", "Indexed mapping"]
    chart.valueAxis.valueMin=0
    chart.valueAxis.valueMax=65
    chart.valueAxis.valueStep=10
    chart.bars[0].fillColor=__import__('reportlab.lib.colors',fromlist=['HexColor']).HexColor('#37627e')
    chart.categoryAxis.labels.fontName='Helvetica'
    chart.categoryAxis.labels.fontSize=13
    chart.valueAxis.labels.fontSize=11
    drawing.add(chart)
    drawing.add(String(360,325,"Median frame processing time in the paired CPU experiment",fontName='Helvetica-Bold',fontSize=15,textAnchor='middle'))
    drawing.add(String(16,185,"ms",fontSize=12))
    drawing.add(String(218,85+oracle/65*230,f"{oracle:.2f} ms",fontSize=14,textAnchor='middle'))
    drawing.add(String(505,85+indexed/65*230,f"{indexed:.2f} ms",fontSize=14,textAnchor='middle'))
    drawing.add(String(360,20,"3 alternating pairs; 60 frames per run; first 6 excluded; same final code and crop",fontSize=10,textAnchor='middle'))
    svg=ASSETS/'paired-latency.svg'
    renderSVG.drawToFile(drawing,str(svg))
    subprocess.run([NODE,'-e',f"require({json.dumps(SHARP)})({json.dumps(str(svg))}).resize(1800).png().toFile({json.dumps(str(svg.with_suffix('.png')))})"],check=True)


doc=Document()
section=doc.sections[0]
section.page_width,section.page_height=Inches(8.2677),Inches(11.6929)
section.top_margin,section.bottom_margin=Inches(.73),Inches(.65)
section.left_margin=section.right_margin=Inches(.70)
section.header_distance=section.footer_distance=Inches(.28)
styles=doc.styles
for style in styles:
    if style.type == 1:
        for border in list(style.element.findall('.//'+qn('w:pBdr'))):
            border.getparent().remove(border)
for name in ('Normal','Body Text'):
    styles[name].font.name='Times New Roman'
    styles[name].font.size=Pt(11)
    styles[name].paragraph_format.line_spacing=1.10
    styles[name].paragraph_format.space_after=Pt(7)
for name,size in [('Title',20),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
    styles[name].font.name='Times New Roman'
    styles[name].font.size=Pt(size)
    styles[name].font.color.rgb=RGBColor(0,0,0)
    styles[name].font.bold=True
    styles[name].paragraph_format.space_before=Pt(8)
    styles[name].paragraph_format.space_after=Pt(7)
styles['Caption'].font.name='Times New Roman'
styles['Caption'].font.size=Pt(9)
styles['Caption'].font.color.rgb=RGBColor(45,45,45)
head=section.header.paragraphs[0]
head.text='MTR-KV     Invention Disclosure Format B     Detailed technical annex'
head.runs[0].font.size=Pt(9)
foot=section.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('Confidential inventor review   |   Annex A')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
for run in foot.runs: run.font.size=Pt(9)
page_titles=[]


def page(title):
    if page_titles: doc.add_page_break()
    page_titles.append(title)
    doc.add_heading(title,level=1)


def p(value, bold=False):
    paragraph=doc.add_paragraph()
    paragraph.add_run(str(value)).bold=bold
    return paragraph


def h(value): doc.add_heading(value,level=2)


def table(headers,rows,widths=None,size=9.3):
    result=doc.add_table(rows=1,cols=len(headers))
    result.alignment=WD_TABLE_ALIGNMENT.CENTER
    result.autofit=False
    if widths:
        for column,width in zip(result.columns,widths):
            column.width=Inches(width)
    for c,value in zip(result.rows[0].cells,headers): c.text=str(value)
    for row in rows:
        for cell,value in zip(result.add_row().cells,row): cell.text=str(value)
    for row_index,row in enumerate(result.rows):
        pr=row._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');pr.append(no_split)
        if row_index==0:
            repeat=OxmlElement('w:tblHeader');pr.append(repeat)
        for col_index,cell in enumerate(row.cells):
            if widths: cell.width=Inches(widths[col_index])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6ECF0' if row_index==0 else ('F5F7F8' if row_index%2==0 else 'FFFFFF'));tcpr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(4)
                paragraph.paragraph_format.space_before=Pt(2)
                paragraph.paragraph_format.line_spacing=1.05
                for run in paragraph.runs:
                    run.font.name='Times New Roman';run.font.size=Pt(size);run.bold=row_index==0
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return result


def picture(name,caption,width=6.8):
    path=ASSETS/(name+'.png') if not name.startswith('prototype') else ROOT/'output/evidence'/(name+'.png')
    doc.add_picture(str(path),width=Inches(width))
    doc.add_paragraph(caption,style='Caption')


def code(lines):
    for line in lines.splitlines():
        q=doc.add_paragraph();q.paragraph_format.space_after=Pt(1)
        q.paragraph_format.line_spacing=1.0
        r=q.add_run(line);r.font.name='Courier New';r.font.size=Pt(8.8)


def link(label,url):
    paragraph=doc.add_paragraph()
    hyperlink=OxmlElement('w:hyperlink')
    rel=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyperlink.set(qn('r:id'),rel)
    run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'174F70');props.append(color)
    run.append(props);txt=OxmlElement('w:t');txt.text=label;run.append(txt);hyperlink.append(run);paragraph._p.append(hyperlink)


def contents():
    page('1 Title of the invention')
    doc.add_paragraph('MTR KV motion guided residency control for visual key value memory',style='Title')
    p('Detailed responses to Invention Disclosure Format IDF-B, Document 02-IPR-R003, Issue 2 dated 01 February 2024. Prepared 21 October 2026 for VIT IPR and TT Cell and inventor review.')
    p('The accompanying PDF preserves the supplied two-page form and adds this annex. This editable companion contains the expanded technical responses. The disclosed executable embodiment is a CPU software prototype: compressed-video metadata controls logical KV-page residency, a generation-safe controller commits mappings, and later attention-derived scores correct placement. A bounded, exact-input-validated reuse path retains final-block ViT K/V across frames.')
    h('2 Field and area of invention')
    p('Computer architecture and systems software for video-transformer inference; compressed-domain video metadata; visual-token geometry; paged key-value memory management; heterogeneous tensor representations; versioned control protocols; and semantic feedback for residency decisions. The implementation demonstrates one final attention block of a ViT-Tiny image classifier, not an all-layer vision-language model.')
    h('Disclosure record and authorship facts')
    p('The frozen architecture dated 22 September 2026 names Kunwar Shaurya Pratap Singh, Kanan Tayal, and Suhani Gupta, B.Tech CSE (AI and Data Engineering), VIT Vellore. These are recorded contributors; this document does not determine legal inventorship or ownership. Inventor contributions, institutional rights, and the final applicant list require confirmation.')
    p('The user reported no known public posting, presentation, submission, or patent filing in this chat on 21 October 2026. That is an inventor-supplied statement, not an independent access-history audit. The architecture date and repository hashes are engineering records, not patent priority dates.')
    table(['Term','Meaning'],[['MTRE','Motion-to-Token Reconciliation Engine'],['MKAD','Motion-Keyed Allocation Descriptor'],['KMC','KV Memory Controller behavioral model'],['SRCE','Semantic Residency Correction Engine'],['T0 / T1 / T2','fp16 / CPU int8 / losslessly compressed fp16 tensor representations']],[1.1,5.65])

    page('3 Prior patents and publications from literature')
    p('Focused review of public primary papers and published patent text, checked 21 October 2026. The comparison identifies overlap and candidate technical distinctions; it is not an exhaustive novelty, inventive-step, claim-construction, or freedom-to-operate opinion. References R1-R13 are listed at the end.')
    table(['Reference and public date','Previously disclosed mechanism','Relationship to this disclosure'],[
        ['R1 PagedAttention / vLLM\n12 Sep 2023','Non-contiguous paged KV memory, block sharing, and copy-on-write.','Paging and logical-to-physical blocks are established. The candidate here adds a codec-originated page-intent and correction protocol.'],
        ['R2 CodecSight v4\nFirst 7 Apr 2026\nRevised 15 Sep 2026','Codec-guided patch pruning, GOP/frame-type refresh, and position-corrected KV reuse across VLM windows.','Closest identified research lead. Broad codec-guided KV reuse is already disclosed. Compare the precise residency transaction and model-feedback correction contract.'],
        ['R3 CoPE-VideoLM v2\nFirst 13 Feb 2026','Motion vectors and residuals encoded into compact video-language representations.','Codec evidence as model input is established. Here that evidence controls residency descriptors; no codec encoder training is implemented.'],
        ['R4 RESUME v1\n30 Sep 2026','An I-frame initializes a latent state; predictive-frame codec observations update it and produce VideoLM readouts.','Carried codec-derived state is related, but is not the tensor-store mechanism implemented here. Its filing-date relevance remains for counsel.'],
        ['R5 H2O v3\nFirst 24 Jun 2023','Attention-heavy and recent tokens guide KV retention.','Attention-based importance is established; it must not be asserted as new by itself.'],
        ['R6 SnapKV v2\nFirst 22 Apr 2024','An observation window selects clustered important KV positions per attention head.','Model-derived cache compression is established. Our software corrects a separate early codec tier proposal through the same versioned controller.'],
    ],[1.55,2.55,2.65],9.1)
    p('All listed research predates this disclosure preparation. RESUME follows the frozen architecture date, but a document date alone does not exclude a publication from a later filing assessment. No publication is excluded here on that basis.')

    page('3 Prior patent comparison continued')
    table(['Published patent text','Disclosed territory reviewed','Technical distinction to examine'],[
        ['R7 US6072548A\nPublished 6 Jun 2000','Dynamic decoder memory allocation for variable decoded-image size.','Decoder-originated memory allocation is old. Our target is transformer visual K/V rather than decoded-picture buffers.'],
        ['R8 US20100098165A1 / US8300697B2\nPublished application 22 Apr 2010','Recursive motion-vector analysis of regions related to decoder memory fetching.','Thresholded regional motion analysis is old. Tiling, hierarchy, or thresholds alone are not novelty anchors.'],
        ['R9 US20250124105A1\n17 Apr 2025\nClaim 1 and description','Binary token vectors are bucketed; representative tokens supply representative cached K/V.','The prototype does not replace token groups with a representative token. It retains page identity and changes representation/residency.'],
        ['R10 US20250225134A1\n10 Jul 2025\nClaim 1 and description','Inference requests reserve contiguous virtual cache space and generate pointers by layer/block.','Reservation and pointer management are established. Compare source metadata, version checks, readiness, and correction relationships.'],
        ['R11 US20260080217A1\n19 Mar 2026\nClaims 1 to 4 and description','Gauge transformations canonicalize attention weights; embodiments split hot-window and cold-tail KV and compress cold data.','This prototype uses unchanged model weights, fp16/int8/raw-byte compression, and motion/semantic page intent. Generic hot/cold KV compression is not claimed as new.'],
    ],[1.65,2.55,2.55],9.3)
    h('Search scope and remaining work')
    p('Searches included codec metadata and KV memory management, motion vectors with semantic cache correction, generation-tagged transformer pages, and attention-guided KV retention. The reviewed pages include paper methods and selected patent independent claims. Published patent text was accessed through Google Patents; no legal-status conclusion is drawn from that mirror.')
    p('A professional search must still trace patent families, earlier priority documents, cited and citing documents, and relevant Indian, PCT, US, EP, and other records. Combining known codec guidance, cache paging, and feedback can still be obvious; a difference in labels or an omission from a summary does not establish inventiveness.')

    page('4 Summary and background of the invention')
    p('Video transformers repeatedly create and retain key/value tensors for visual tokens. Compressed video exposes spatial change before the model has produced semantic attention. Using that early signal can reduce fast-tier pressure, but low motion is not the same as low importance: a stationary object, text region, or scene cue may still matter to the task. A second problem is consistency. A delayed descriptor or correction must not install stale data after a newer frame has taken ownership of a page.')
    p('MTR-KV separates early evidence from later verification. A decoder adapter produces source-coordinate blocks, motion magnitude, confidence, and frame type. MTRE reconciles those blocks with the model preprocessing and token/page layout. A temporal engine and scene guard qualify the proposal. MKAD records the stream, generation, page span, operation, and target tier. KMC validates and commits a complete mapping before issuing readiness. Actual tensor representations are then consumed by attention.')
    p('After model execution, class-token attention is aggregated into page scores. SRCE identifies disagreement, particularly low codec activity with high model importance, and sends a generation-matched correction through the same descriptor/controller path. The software never bypasses generation checks for a semantic rescue. Missing metadata takes a conservative branch instead of being treated as a zero-motion observation.')
    h('Candidate distinguishing technical contribution')
    p('The proposed contribution is the coordinated control relationship: source-space codec evidence becomes page-specific tier intent before attention; a versioned reserve/prepare/commit/ready sequence gates use; and later model-derived importance corrects that intent through the same controller while stale generations are rejected. The reuse extension requires both codec/temporal eligibility and exact final-block page-input equality, and binds retained content to the new generation only after readiness.')
    p('This combination is the candidate for examination, not a conclusion that no earlier system implements it. R2 already establishes codec-guided VLM cache refresh, while R1 and R5-R11 cover adjacent memory and importance mechanisms. The claimable distinction, if any, must lie in concrete relationships and enabling details rather than the names MTRE, MKAD, KMC, or SRCE.')
    h('What the prototype contributes today')
    p('The implemented artifact provides executable control semantics, a real compressed-video path, actual tensor representations, attention consumption, bounded reuse with invalidation, traceable correction actions, an indexed reconciliation path checked against an oracle, and a local visual demonstration. It does not provide custom hardware, all-layer VLM state, task-labeled semantic correctness, or proof of commercial real-time performance.')

    page('4 Candidate distinction matrix')
    p('The matrix separates an implemented local mechanism from the strongest known overlap. “Compare” identifies a concrete issue to investigate; it does not mean an earlier reference lacks the feature.')
    table(['Element in MTR-KV','Relevant overlap','Candidate relationship to examine'],[
        ['Actual inverse preprocessing maps codec blocks to token pages.','R2 maps patch selection through model layout; R3 uses codec primitives.','Preserve source/crop/page identity through generation-tagged residency descriptors, including partial blocks and missing metadata.'],
        ['Motion and temporal state propose T0/T1/T2 intent before attention.','Codec-guided decisions in R2; hot/cold KV in R11.','Tie a codec-derived spatial proposal to a page-tier operation rather than visual pruning or solely recency-based placement.'],
        ['Reserve, prepare, commit, then ready; reject stale descriptors.','Paging/reservation in R1 and R10; general versioned transactions are established systems techniques.','Define the exact generation/range/operation checks, atomic publication model, failure preservation, and completion contract for this pipeline.'],
        ['Model importance rescues a quiet page through the same KMC.','Importance-driven retention in R5/R6; R2 also offers residual-based rescue.','Use later model attention, not codec residual alone, to correct the earlier residency intent without bypassing the consistency path.'],
        ['Static-page reuse also checks exact model input, owner, age, and versions.','R2 uses GOP/frame identity and position correction for approximate VLM reuse.','The narrow final-block implementation avoids reuse solely on motion assumptions and explicitly rebinds retained content after new-generation readiness.'],
        ['Conservative invalidation and auditable trace.','Fallback, validation, and tracing are established engineering practices.','Specify their interaction with the proposed controller arrangement; do not claim these practices individually.'],
    ],[2.1,2.2,2.45],9.4)
    p('The strongest present position is a narrowly described technical arrangement supported by code and tests. The weakest positions would be “motion vectors are new,” “paging is new,” “attention rescue is new,” or “cross-frame KV reuse is new.” These are not asserted.')
    p('The newly measured optimization uses conventional spatial indexing and compact tensor serialization. Those improvements make the implementation more practical, but this disclosure does not rely on them as independent novelty claims.')

    page('5 Objectives of invention')
    for value in [
        'Use early compressed-video evidence to propose visual-page residency before the relevant model attention computation.',
        'Reconcile actual source pixels, preprocessing crop/resize, patch tokens, and logical page spans without assuming codec blocks and patches align.',
        'Preserve generation consistency during allocation, replacement, semantic correction, and cross-frame payload reuse.',
        'Retain conservative behavior when evidence is missing, uncertain, or transitioning; avoid treating absence of metadata as proof of static content.',
        'Recover important low-motion pages using later model-derived scores through the normal controller transaction.',
        'Measure real tensor payloads and actual execution, distinguish them from simulator capacity estimates, and expose reproducible evidence.',
    ]: p(value)
    h('6 Working principle of the invention in brief')
    p('Decode a frame and extract motion-vector side data. Normalize it, map source regions to the actual model token grid, and aggregate activity by logical page. Update temporal and transition state. Construct one or more generation-tagged MKADs, validate them, reserve replacement space, prepare the complete shadow mapping, and publish it before emitting a ready token.')
    p('The model can then attach or reuse K/V for the current ready generation. A reuse candidate must also pass exact feature, owner, parameter, and age checks. Prefix tokens and queries are computed normally. Attention reads every patch page through the tier store. After attention, SRCE compares page importance with early motion evidence and requests safe corrections through KMC. The next frame carries temporal state forward; old-generation updates remain stale.')
    p('The live implementation is sequential around the Python controller and model. The event simulator separately models asynchronous ordering and adversarial delays. Physical overlap between a media engine and GPU/RTL controller is a proposed embodiment, not a measured property of the Python run.')

    page('7 Detailed description of the invention')
    picture('architecture','Figure 1. Implemented software control/data path. The dotted hardware/VLM extension is proposed. KMC readiness precedes model access; the feedback loop returns through the same controller.')
    h('Module contracts and data direction')
    p('The decoder adapter supplies blocks, vector scale, source dimensions, presentation timestamp, and frame type. MotionNormalizer creates canonical MotionBlock records. MTRE produces token and page activity; temporal and scene logic govern conservative residency. MkadBuilder groups compatible adjacent page ranges into MKADs. KMC owns the authoritative logical mapping and readiness state.')
    p('TorchKVStore holds actual model tensors associated with the committed page generation. ViTAttentionFeedback replaces only the last attention block in the model-backed path. SRCE produces corrections from later class-token attention. JSONL records and the dashboard expose measurements but do not participate in the allocation decision.')
    p('The live controller processes one stream per instance. The simulator also includes a multistream workload. The phone adapter has one exclusive capture session per server; it decodes continuously and retains the latest frame when model work is slower than capture.')

    page('7 Source geometry and spatial reconciliation')
    picture('geometry','Figure 2. Source-to-token mapping for the measured 640 x 360 source and selected ViT transform. The model sees a centered crop; colored token groups illustrate pages, not a spatial heat map.')
    p('Let a model patch occupy rectangle T in source coordinates after inverting preprocessing. For codec block B, overlap weight is area(B intersect T) divided by area(T). Token evidence is accumulated with these weights. Mixed active/static overlap is treated conservatively: a high-change component can dominate the plain mean. Confidence remains an independent guard.')
    p('The optimized path indexes blocks into spatial tiles, enumerates only candidates that can intersect each patch, and restores original block order before accumulation. Global page aggregation is then performed once; it is not incorrectly repeated as one page per token. Randomized mixed-block tests, duplicate observations, boundary tiles, and persistence values match the complete oracle exactly.')
    p('The previous live mapping assumed a full-frame stretch while the selected model used resize and center crop. This build fixes that mismatch by deriving the source crop from the actual timm transform. Geometry tests check the resulting bounds. Unsupported padded cropping is rejected rather than silently assigned incorrect page coordinates.')

    page('7 Generation safe transaction and descriptor')
    picture('transaction','Figure 3. Behavioral transaction and failure preservation. Atomic publication and the fence are modeled in Python; this figure does not assert hardware atomicity or a physical asynchronous FIFO.')
    table(['Descriptor field','Role'],[
        ['stream_id and generation','Identify the source stream and ownership version.'],
        ['page_start and page_count','Specify a contiguous logical page range; range is validated.'],
        ['operation and target_tier','Allocate, promote, demote, retire, or prefetch intent.'],
        ['activity, confidence, cause','Record evidence and the reason for the operation.'],
        ['descriptor_id and frame_index','Support idempotence and traceability. IDs advance for both normal and correction descriptors.'],
    ],[2.0,4.75])
    p('A successful allocation first reserves all required replacement blocks. The old mapping remains active while the shadow mapping is prepared. Publication updates the complete page range and committed generation. A ready token records stream, generation, span, and modeled fence ID. Tensor access checks both readiness and exact generation.')
    p('A failed replacement reservation preserves the old mapping; fallback or abort is reported without pretending the new generation is ready. Duplicate descriptor handling and stale rejection are separately tested. The software allocator uses a shared pool with replacement headroom, not physically distinct tier banks.')

    page('7 Bounded cross frame reuse')
    picture('reuse','Figure 4. Reuse gates. Readiness failure blocks model access; a later failed reuse gate refreshes the page. Passing candidates rebind retained payloads only after current-generation readiness.',width=6.5)
    p('The exact-input check applies at the final attention block, after earlier global attention. This matters because low pixel motion does not guarantee unchanged contextual features. The adapter computes Q for every token and K/V only for prefix tokens and misses; attention still reads each patch page through the store.')
    p('The default maximum origin age is eight processed generations. Tier migration preserves origin_generation, so repeated promotion cannot indefinitely renew stale content. Only the latest validation features and payloads are retained. Inference failure, controller change, generation gaps, and parameter changes invalidate history.')

    page('7 Tensor representations and semantic correction')
    table(['Tier','Actual implementation','Boundary'],[
        ['T0 HOT','Contiguous fp16 tensors on the selected torch device.','The verified runs use CPU; a selected torch device is not a custom GPU controller.'],
        ['T1 WARM','Symmetric per-tensor int8 on CPU with a scale, decoded on read.','Lossy relative to original fp16; classifier/logit comparison measures observed deviation.'],
        ['T2 COLD','Shape header plus zlib-compressed contiguous fp16 bytes on CPU.','Lossless relative to the fp16 payload. Archive/pickle overhead was removed in this build.'],
        ['T3','Optional RECONSTRUCT controller/simulator state.','Actual tensor reconstruction is not implemented; tensor writes to T3 are rejected.'],
    ],[1.0,3.0,2.75])
    p('The tensor store records payload bytes written/read and synchronized operation time. Payload accounting excludes Python object overhead, model weights, decoder buffers, and reuse-validation features. The controller separately estimates capacity using its configured tier ratios; that estimate is not an actual measured tensor allocation.')
    p('Class-token attention from the final ViT block is averaged over heads, normalized, and aggregated by page. A low-motion page with high score is a candidate for SRCE promotion. This is a model signal, not a human-labeled measure of relevance. The score arrives after that frame has completed inference, so it does not avoid the pass that produced it.')
    p('Corrections carry the current generation through MkadBuilder and KMC. The store synchronizes only representations whose tier actually changed. Unchanged pages are not decoded again merely to refresh a snapshot. Actual reads required inside attention remain counted.')
    h('Illustrative failure sequence')
    p('If generation 12 is current and a semantic correction for generation 11 arrives, the live controller returns STALE_DROP. A current-generation quiet page can be promoted; its retained tensor remains associated with the current generation and original content age. A subsequent attempt to read generation 11 fails even when the underlying content had been reused.')

    page('7 Runtime inputs and implementation scope')
    p('File input uses PyAV/FFmpeg export_mvs and emits one JSONL record per processed frame. A missing-vector intra frame, unsupported codec, or unreliable metadata enters fallback. Residual energy is unavailable through this adapter and is not fabricated. FFmpeg-compatible encoded network inputs are accepted, subject to decoder support.')
    p('The Android USB mode uses scrcpy direct-camera H.264 capture and ADB on Android 12 or newer. A private FIFO carries Matroska bytes without saving a camera recording, audio is disabled, and the decoder keeps one latest decoded frame. Software tests exercise a real encoded producer, timeout/cleanup, and exclusive sessions. scrcpy/ADB installation and a physical-phone trial remain pending; neither is presented as completed evidence.')
    h('Optimizations implemented in this build')
    table(['Change','Reason and verification'],[
        ['Enable tiled reconciliation in LiveController','Avoid all-block scans per patch. Restore original accumulation order and compare complete token/page outputs against the oracle.'],
        ['Correct page aggregation','Combine all token records before page aggregation; prevent duplicate per-token page summaries.'],
        ['Align actual preprocessing geometry','Invert the selected model resize/center crop and apply it to token regions. Bounds are covered by tests.'],
        ['Use compact T2 bytes','Compress fp16 bytes and shape metadata without torch.save ZIP/pickle containers. Scalar, empty, and multidimensional round trips are exact after fp16 conversion.'],
        ['Synchronize only changed tiers','Remove unnecessary post-attention full-page rereads. Migration I/O remains measured.'],
    ],[2.1,4.65])
    p('The correctness oracle remains selectable with mtre.use_optimized_path=false in configuration. The frozen architecture document is preserved. The default implementation path changes, but the candidate transaction and correction sequence remains the same.')
    p('Not implemented: a physical GPU/RTL KMC, an all-layer VLM integration, cross-window language-model context management, production asynchronous media/GPU overlap, generic browser/webcam upload, or a representative downstream task benchmark.')

    page('8 Experimental validation method')
    p('Evidence was generated from the current local working tree on 21 October 2026. It is a self-run engineering validation, not an independent replication. The generated H.264 source is 640 x 360, 12 fps, and 216 frames, with still, moving-object, still, and scene-cut phases. Generated input is deliberately labeled; it is not described as a representative external dataset.')
    table(['Setting','Value'],[
        ['Model','vit_tiny_patch16_224.augreg_in21k_ft_in1k, pretrained ViT-Tiny image classifier'],
        ['Scope','Final attention block; 196 patch tokens, 32 tokens/page, 7 pages; prefix tokens remain in the model'],
        ['Device','CPU; this runtime reports CUDA unavailable and MPS unavailable'],
        ['Paired experiment','3 alternating oracle/indexed pairs; 60 frames/run; first 6 excluded; identical crop, storage, model, source, and reuse settings'],
        ['Comparison inside each run','Standard model inference on the same frame plus stored-K/V inference; top-1 and mean absolute logit error recorded'],
        ['Timing boundary','Decoder plus processing; excludes model loading, JSON serialization, dashboard rendering, and presentation pacing'],
        ['Functional run','A separate serial 216-frame run covering all four generated phases'],
    ],[1.65,5.1])
    picture('paired-latency','Figure 5. Primary timing comparison: the indexed and reference reconciliation paths in the same optimized build. This isolates mapping choice; it is not a comparison against a full production VLM.',width=6.6)

    page('8 Measured results and interpretation')
    rows=[]
    for run in paired['runs']:
        summary=run['summary']
        rows.append([str(run['repeat']),run['arm'],str(summary['frames']),f"{summary['mean_end_to_end_frame_ms']:.3f}",f"{summary['p95_end_to_end_frame_ms']:.3f}",str(summary['kv_pages_reused']),f"{100*summary['top1_agreement_rate']:.1f}%"])
    table(['Pair','Path','Frames','Mean ms','P95 ms','Page reuses','Top 1 agreement'],rows,[.45,.75,.65,.75,.75,1.2,1.2],9.2)
    p(f'The median of the three per-run mean frame times is {oracle:.2f} ms for the oracle and {indexed:.2f} ms for indexed mapping. The observed ratio is {paired["observed_pipeline_ratio"]:.2f}x, equivalent to {100*(1-indexed/oracle):.2f}% lower frame time in this experiment. All paired controller-decision hashes match, including page tiers, descriptors, readiness, correction actions, and source crop.')
    h('Full serial generated video run')
    table(['Measurement','Observed value'],[
        ['Frames ready / compared',f'{full["ready_frames"]} / {full["compared_frames"]}'],
        ['Page reuses / recomputed pages',f'{full["kv_pages_reused"]} / {full["kv_pages_recomputed"]}'],
        ['Patch K/V tokens reused',str(full['kv_tokens_reused'])],
        ['Top 1 agreement / mean absolute logit error',f'{100*full["top1_agreement_rate"]:.1f}% / {full["mean_abs_logit_error"]:.8f}'],
        ['Mean / P95 decoder plus processing time',f'{full["mean_end_to_end_frame_ms"]:.3f} / {full["p95_end_to_end_frame_ms"]:.3f} ms'],
        ['Metadata fallback',f'{full["metadata_fallback_rate"]*100:.3f}% of frames'],
        ['Peak tensor payload / peak validation input buffers',f'{full["peak_actual_tensor_payload_bytes"]:,} / {full["peak_reuse_validation_input_bytes"]:,} bytes'],
    ],[3.5,3.25])
    p('These results demonstrate working reuse and an indexed-mapping improvement on this sample. They do not show a memory reduction for the whole model: validation buffers are additional, and the payload peak excludes model weights. They also do not establish a speedup over standard inference. In the full run, the baseline model and stored-K/V model times are separately recorded; earlier ViT blocks still execute.')
    p('The experiment has no task labels, confidence interval, thermal controls, power measurement, or external corpus. Top-1 agreement is only classifier consistency on these frames. Representative static/motion/scene-change videos and downstream quality tasks remain necessary before broader numerical assertions.')

    page('8 Prototype screenshot of video and page residency')
    picture('prototype-video-and-pages','Figure 6. Authentic local dashboard screenshot from the optimized build at frame 48. It shows the decoded source, seven committed logical pages, three T0 pages after semantic correction, seven K/V page reuses, and zero refreshed pages for this observation.')
    p('This screenshot is an actual UI capture, not a mockup. The displayed 33.22 fps is the inverse of that frame\'s decoder/processing time; it excludes rendering and pacing and is not camera-to-display latency. The run completion time includes initialization and presentation effects and should not be compared directly with the serial command-line benchmark.')
    p('The colored page tiles follow logical token-page order. They do not draw rectangular source regions or identify objects. The source preview displays the full decoded frame while the model consumes the explicitly recorded center crop. Trace fields preserve that geometry.')
    h('How to inspect the same behavior')
    code('.venv/bin/python -m mtr_kv dashboard results/demo.mp4 \\\n  --semantic-model --kv-inference --kv-reuse \\\n  --compare-baseline --max-frames 48 --port 8768')
    p('Open http://127.0.0.1:8768 and restart the run. Hover or focus a page tile to inspect its temporal state and activity. Load a simulator artifact separately only when comparing simulated strategies; it is not the source of the measured live cards.')

    page('8 Prototype screenshot of actual reuse measurements')
    picture('prototype-reuse-trace','Figure 7. Authentic measured-frame table. Frame 43 refreshes all pages at the age bound; nearby static frames reuse seven pages. The rows include actual tensor payloads, model time, and standard-model top-1 comparison.')
    p('The screenshot shows 143,186 bytes of stored K/V at these observations. This differs from the run-wide peak because page representations change. The comparison result belongs to the displayed generated input and device. It does not demonstrate semantic understanding, object detection, or a VLM question-answering score.')
    p('The screenshot files are preserved in the evidence bundle with hashes. The primary machine-readable proof is the raw JSONL, not the screenshot alone. Raw records include generation, source timestamps, extracted vector count, descriptors, apply statuses, ready tokens, page states, correction actions, tensor I/O, reuse provenance, and baseline comparisons.')
    h('Observable controls')
    p('Restart begins a new controller session; Stop closes the stream; Save trace exports image-free JSONL. The table retains a bounded recent display. Complete traces for this disclosure were captured through the CLI, so they are not limited by the browser retention window.')

    page('8 Test evidence and reproducibility record')
    p('The optimized build passed 128 automated tests. The saved text report and JUnit XML record 128 passes and no failures. Tests cover controller behavior, adversarial ordering, semantic correction, geometry, actual tensor representations, model consumption, temporal reuse, encoded video, local HTTP streaming, and the simulated USB camera transport.')
    table(['Safety or correctness case','Evidence entry point'],[
        ['Stale ownership, duplicate descriptors, reservation failure preservation','tests/test_memory_generation.py and tests/test_race.py'],
        ['Pixel crop bounds and exact indexed/oracle output equality','tests/test_geometry.py'],
        ['Quiet codec evidence with changed final-block features','tests/test_kv_reuse.py'],
        ['Intra frames, missing metadata, generation gaps, age expiry','tests/test_kv_reuse.py and tests/test_live.py'],
        ['Actual fp16/int8/compressed tensor round trips and tier migration','tests/test_tensor_store.py'],
        ['Full model/decoder reuse and observed classifier comparison','tests/test_video.py'],
        ['USB transport cleanup and exclusive browser sessions','tests/test_phone.py and tests/test_dashboard_server.py'],
    ],[3.5,3.25])
    p('The randomized geometry regression includes mixed valid/invalid blocks, out-of-bounds rectangles, different tile counts, duplicate observations, page persistence, and full token/page equality. Raw compressed fp16 tests include scalar, empty, two-dimensional, and three-dimensional tensors. These checks support software behavior; they do not prove physical-device atomicity or formal RTL correctness.')
    h('Evidence package contents')
    p('MTR_KV_Evidence.zip contains the generated source video, full serial trace and evaluation, all six paired traces and their manifest, test text/XML, screenshots, the original supplied form, source/configuration snapshots, references and search record, and a SHA-256 manifest. A verification script checks hashes. Hashes identify local bytes and changes; they are not a trusted timestamp or independent certification.')
    p('The code snapshot includes uncommitted Python modules because the working tree was not clean. The package therefore preserves exact source bytes rather than implying that Git HEAD alone identifies this build. Environment versions and the dirty-tree state are recorded. The original frozen architecture remains a separate design record.')

    page('9 Aspects of the invention for which protection is sought')
    p('Protection is requested for the specific coordinated controller arrangement, subject to professional prior-art and enablement review. The following are technical topics for claim drafting, not final legal claims.')
    h('Primary coordinated mechanism')
    p('A processing system receives compressed-video change metadata before the relevant model attention; reconciles source regions with actual visual-token and page geometry; uses temporal and confidence state to issue a generation-tagged page-range and tier descriptor; validates, reserves, prepares, commits, and signals readiness before use; and applies later model-derived page-importance correction through the same generation-checked controller path.')
    h('Supporting consistency and reuse relationships')
    p('A method preserves the active mapping on failed replacement reservation, rejects stale or duplicate operations, and ties access to a current-generation readiness result. A temporal reuse method retains actual page K/V only when codec/temporal eligibility and exact model-input identity both hold, checks controller/model/geometry ownership and an origin-age bound, and rebinds retained content after the replacement generation is ready. Otherwise the page is recomputed.')
    h('Potential dependent technical details')
    p('Candidate dependent topics include inverse crop/resize mapping; conservative mixed-evidence aggregation; missing-metadata fallback; distinct confidence and activity signals; current-generation low-motion/high-attention promotion; page-span coalescing; explicit source and origin generations; and representation migration without age renewal. The implemented descriptors, predicates, failure cases, and measurements provide details for narrower drafting.')
    h('What should not be asserted as new alone')
    p('Motion vectors, thresholds, spatial indexing, paging, quantization, compression, attention-derived importance, generic transactions, and KV reuse are established techniques. The Android camera adapter, dashboard, compact serialization, and faster Python mapping are implementation conveniences, not the principal proposed invention. The searches reviewed here do not establish that the proposed combination is absent from earlier art or non-obvious.')
    p('Hardware media engines, asynchronous FIFOs, physical tier banks, GPU page-table publication, and all-layer VLM windows are proposed extensions. Their claim breadth requires a separately enabling design and inventor confirmation. A Python fence counter must not be described as a fabricated or tested hardware fence.')

    page('10 Technology readiness level')
    p('Selected overall level: TRL 3 - experimental proof of concept. This is a conservative self-assessment for the complete disclosed architecture, based on the supplied template scale. The software embodiment has been exercised in a local lab runtime, but the broader hardware/VLM system has not been validated in a relevant deployment environment.',bold=True)
    table(['TRL','Template meaning','Assessment'],[
        ['1','Basic principles observed','Exceeded for the software concept.'],
        ['2','Technology concept formulated','Documented architecture and executable interfaces exist.'],
        ['3','Experimental proof of concept','SELECTED: functioning software, real encoded input, tensor consumption, tests, and measurements.'],
        ['4','Technology validated in a lab','Some software components have local validation; no overall hardware/VLM qualification is asserted.'],
        ['5 to 6','Validated or demonstrated in a relevant environment','Not established: representative tasks, device integration, and operational constraints remain untested.'],
        ['7 to 9','Operational prototype, qualified system, or proven deployment','Not established.'],
    ],[.55,2.6,3.6])
    h('Work required for the next readiness assessment')
    p('Collect licensed representative videos and task labels, preserve input provenance, compare codec-only/semantic/reuse variants, record repeated latency and actual device memory, and review errors rather than only agreement. Test the physical Android phone after dependency installation. Select a target VLM and hardware runtime before claiming all-layer or device-level benefit.')
    h('Inventor and institutional confirmation before submission')
    p('Confirm inventorship contributions, ownership under VIT policy and any sponsorship/employment, conception evidence, unpublished design records, third-party code/model/data licenses, and the absence of earlier public disclosures. The user supplied no known disclosure or filing; this remains subject to confirmation. A patent agent should review claim scope and prior art before the form is treated as a filing specification. Section 10 requirements and disclosure timing are discussed in R12-R13.')
    p('This completed IDF-B is a technical invention-disclosure record for review. It is not a patent-office filing, grant prediction, legal novelty opinion, or institutional approval. No signature, applicant identity, or legal inventorship declaration has been fabricated.')

    page('References and reviewed source locations')
    refs=[
      ('R1','Kwon et al. Efficient Memory Management for Large Language Model Serving with PagedAttention. arXiv:2309.06180v1, 12 Sep 2023; SOSP 2023. Methods sections 4.1-4.4.','https://arxiv.org/html/2309.06180v1'),
      ('R2','Zou et al. CodecSight Leveraging Video Codec Signals for Efficient Streaming VLM Inference. arXiv:2604.06036v4, 15 Sep 2026; first submitted 7 Apr 2026. Sections 3.3-3.5.','https://arxiv.org/html/2604.06036v4'),
      ('R3','Sarkar et al. CoPE-VideoLM Leveraging Codec Primitives For Efficient Video Language Modeling. arXiv:2602.13191v2, 30 Mar 2026; first submitted 13 Feb 2026. Section 3.','https://arxiv.org/html/2602.13191v2'),
      ('R4','Zhang et al. RESUME Recurrent State Updates from Motion and Residual Signals for Efficient Video Language Modeling. arXiv:2609.39563v1, 30 Sep 2026. Section 3.','https://arxiv.org/html/2609.39563v1'),
      ('R5','Zhang et al. H2O Heavy-Hitter Oracle for Efficient Generative Inference of Large Language Models. arXiv:2306.14048v3; first submitted 24 Jun 2023.','https://arxiv.org/abs/2306.14048'),
      ('R6','Li et al. SnapKV LLM Knows What You are Looking for Before Generation. arXiv:2404.14469v2; first submitted 22 Apr 2024.','https://arxiv.org/abs/2404.14469'),
    ]
    for key,desc,url in refs:
        p(key+'  '+desc)
        link(url,url)
    p('Source access date: 21 October 2026. Descriptions are paraphrases. Source performance numbers are not transferred to MTR-KV and are not used as proof of this prototype\'s performance.')

    page('References continued and search record')
    refs=[
      ('R7','US6072548A. Video decoder dynamic memory allocation system and method allowing variable decoded image size. Published 6 Jun 2000.','https://patents.google.com/patent/US6072548A/en'),
      ('R8','US20100098165A1 / US8300697B2. Dynamic motion vector analysis method. Published application 22 Apr 2010; grant publication 30 Oct 2012.','https://patents.google.com/patent/US8300697B2/en'),
      ('R9','US20250124105A1. Efficient token pruning in transformer-based neural networks. Published 17 Apr 2025. Reviewed claim 1 and related description.','https://patents.google.com/patent/US20250124105A1/en'),
      ('R10','US20250225134A1. Key-value cache management system for inference processes. Published 10 Jul 2025. Reviewed claim 1 and pointer/reservation description.','https://patents.google.com/patent/US20250225134A1/en'),
      ('R11','US20260080217A1. Key-value cache compression based on gauge transformation. Published 19 Mar 2026. Reviewed claims 1-4 and hot/cold description.','https://patents.google.com/patent/US20260080217A1/en'),
      ('R12','Indian Patents Act 1970, Section 10. Contents of specifications. Official IP India text, accessed 21 Oct 2026.','https://ipindia.gov.in/acts/patent-act-1970/section-10'),
      ('R13','WIPO. Frequently Asked Questions on Patents, including novelty and public-disclosure timing. Accessed 21 Oct 2026.','https://www.wipo.int/en/web/patents/faq_patents'),
    ]
    for key,desc,url in refs:
        p(key+'  '+desc)
        link(url,url)
    p('Search expressions: motion vector transformer KV cache page table; codec metadata descriptor memory controller attention; video decoder semantic cache correction generation; visual token page residency motion residual; KV cache generation motion controller; and attention-guided retention. No-result or broad-result queries are not treated as evidence that an element is absent from prior art.')

    page('Evidence locations and reproduction')
    table(['Local evidence','Purpose'],[
        ['paired-serial/manifest.json','Input hash, settings, six trace hashes, per-run metrics, decision hashes, and median ratio.'],
        ['paired-serial/repeat-N-PATH.jsonl','All raw measured frames from each paired run.'],
        ['full-demo-serial.jsonl and evaluation JSON','Complete generated four-phase run and machine summary.'],
        ['test-results.txt and test-results.xml','128 passing tests and machine-readable test outcomes.'],
        ['prototype-video-and-pages.png and prototype-reuse-trace.png','Unmodified screenshots captured from the live local browser.'],
        ['source snapshot and manifest.json','Exact Python, tests, configuration, frontend, requirements, and SHA-256 inventory.'],
    ],[2.8,3.95],9.5)
    p('Measurements live under results/optimization-20261021 in the working project and are copied into the evidence archive. The verify_evidence.py script checks every listed SHA-256 digest after extraction. No network upload is required to inspect the prototype or evidence.')
    h('Commands used for the optimized proof')
    code('PYTHONPATH=. .venv/bin/python experiments/benchmark_optimization.py \\\n  --output results/optimization-20261021/paired-serial\n\n.venv/bin/python -m mtr_kv video results/demo.mp4 \\\n  --semantic-model --kv-inference --kv-reuse \\\n  --compare-baseline --max-frames 216 > full-demo-serial.jsonl\n\n.venv/bin/python -m pytest -q \\\n  --junitxml=results/optimization-20261021/test-results.xml')
    h('Important implementation entry points')
    p('mtr_kv/mtre.py: reconcile and reconcile_optimized; mtr_kv/live.py: frame, semantic, and reuse_kv; mtr_kv/controller.py: generation checking and allocation; mtr_kv/tensor_store.py: put/get, sync_tiers, and reuse; mtr_kv/vision_feedback.py: source_crop and _reuse_attention_kv; mtr_kv/video.py: decoder/model integration. These are supplied in the code snapshot.')
    p('The generated demonstration is useful for verifying plumbing and controlled cases. Any later claim about production memory pressure, VLM task accuracy, energy, or device throughput needs additional experiments with identified hardware and representative permitted data.')


if __name__=='__main__':
    diagrams()
    contents()
    OUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUT)
    (ROOT/'tmp/disclosure/page-titles.json').write_text(json.dumps(page_titles,indent=2))
    print(OUT)
    print('Planned annex pages',len(page_titles))
