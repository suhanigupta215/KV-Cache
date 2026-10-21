"""Reorganize verified MTR-KV content using the supplied IDF-B reference components."""
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import io
import json
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import build_disclosure as old

ROOT=Path(__file__).resolve().parents[1]
REF=Path('/Users/shaurya/Downloads/Aura_Patent.docx')
TMP=ROOT/'tmp/aura-reference'
ASSETS=TMP/'assets'
OUT=ROOT/'output/docx/MTR_KV_Patent_Disclosure_Revised.docx'
ASSETS.mkdir(exist_ok=True)
contract=(TMP/'artifact.md').read_text()
assert sha256(REF.read_bytes()).hexdigest() in contract
groups=[]
def capture_page(title): groups.append({'title':title,'items':[]})
def cap(kind,*args,**kwargs): groups[-1]['items'].append((kind,args,kwargs))
old.page=capture_page
for name in ['p','h','table','picture','code','link']:
    setattr(old,name,lambda *a,_name=name,**k:cap(_name,*a,**k))
old.contents()
trace_path=ROOT/'results/optimization-20261021/full-demo-serial.jsonl'
trace=[json.loads(line) for line in trace_path.read_text().splitlines()]
by_frame={x['generation']:x for x in trace}
selected=[1,3,8,43,48,55,109,163,216]
(TMP/'implementation-records.json').write_text(json.dumps({
    'source':str(trace_path.relative_to(ROOT)), 'sha256':sha256(trace_path.read_bytes()).hexdigest(),
    'records':[by_frame[n] for n in selected]},indent=2))

# Source components are copied in place; package assembly below preserves opaque parts.
doc=Document(REF)
protos={name:deepcopy(doc.paragraphs[index]._p) for name,index in
        [('question',1),('title',2),('body',5),('sub',33),('caption',128)]}
first_table=deepcopy(doc.tables[0]._tbl)
last_table=deepcopy(doc.tables[7]._tbl)
trl_table=deepcopy(doc.tables[8]._tbl)
break_proto=next(deepcopy(p._p) for p in doc.paragraphs if p._p.find('./'+qn('w:pPr')+'/'+qn('w:sectPr')) is not None)
body=doc._element.body
for child in list(body):
    if child.tag!=qn('w:sectPr'): body.remove(child)
body.insert(0,first_table)
title_style=doc.styles.add_style('Title',WD_STYLE_TYPE.PARAGRAPH)
title_style.base_style=doc.styles['Body Text']
title_style.font.name='Times New Roman'
title_style.font.size=Pt(13)
title_style.font.bold=True
title_style.font.color.rgb=RGBColor(0,0,0)
title_style.paragraph_format.space_after=Pt(6)
doc.core_properties.title='MTR KV Invention Disclosure Format B'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
doc.core_properties.subject='Generation safe motion guided visual KV residency and semantic correction'
doc.core_properties.comments=''

def para(text,role='body',bold=False):
    element=deepcopy(protos[role])
    props=element.find(qn('w:pPr'))
    for child in list(element):
        if child is not props: element.remove(child)
    for prop in ['numPr','sectPr','pageBreakBefore']:
        for child in props.findall(qn('w:'+prop)): props.remove(child)
    if role in ['title','question']:
        ind=props.find(qn('w:ind'))
        if ind is not None: props.remove(ind)
    # Explicit section numbers replace the reference's automatic question numbering.
    if role=='title': props.find(qn('w:pStyle')).set(qn('w:val'),'Title')
    body.insert(len(body)-1,element)
    from docx.text.paragraph import Paragraph
    p=Paragraph(element,doc._body)
    r=p.add_run(str(text));r.font.name='Times New Roman';r.font.size=Pt(13 if role=='title' else 11)
    r.font.color.rgb=RGBColor(0,0,0)
    r.bold=bold or role in ['question','title','sub','caption']
    if role=='body': p.paragraph_format.keep_with_next=False
    return p

def heading(text): return para(text,'sub')
def question(text): return para(text,'question')
def newpage():
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0)
    p.add_run().add_break(__import__('docx.enum.text',fromlist=['WD_BREAK']).WD_BREAK.PAGE)

def hyperlink(p,label,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
    r=OxmlElement('w:r');props=OxmlElement('w:rPr')
    font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Times New Roman');font.set(qn('w:hAnsi'),'Times New Roman');props.append(font)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'20');props.append(size)
    r.append(props);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def table(headers,rows,widths=None,size=10):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    widths=widths or [7.1/len(headers)]*len(headers)
    factor=7.1/sum(widths)
    for col,w in zip(t.columns,widths): col.width=Inches(w*factor)
    for c,text in zip(t.rows[0].cells,headers): c.text=str(text)
    for row in rows:
        for c,text in zip(t.add_row().cells,row): c.text=str(text)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0: pr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]*factor);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            cp.append(borders)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAEAEA' if i==0 else 'FFFFFF');cp.append(sh)
            margin=OxmlElement('w:tcMar')
            for side,value in [('top',80),('bottom',80),('left',90),('right',90)]:
                m=OxmlElement('w:'+side);m.set(qn('w:w'),str(value));m.set(qn('w:type'),'dxa');margin.append(m)
            cp.append(margin)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.04
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=i==0;r.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def picture(name,title,explanation,source=None,width=7):
    heading(title)
    path=ASSETS/(name+'.png')
    if not path.exists():path=ROOT/('output/evidence' if name.startswith('prototype') else 'tmp/disclosure/assets')/(name+'.png')
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(5)
    p.add_run().add_picture(str(path),width=Inches(width))
    para(explanation)
    if source:
        p=para('Source: '+source)
        for r in p.runs:r.font.size=Pt(10)

def emit(group,skip=(),rename=None):
    for idx,(kind,args,kwargs) in enumerate(groups[group]['items']):
        if idx in skip:continue
        if kind=='p':para(args[0],bold=kwargs.get('bold',False))
        elif kind=='h':heading((rename or {}).get(args[0],args[0]))
        elif kind=='table':table(*args,**{k:v for k,v in kwargs.items() if k!='size'})
        elif kind=='link':hyperlink(para(''),*args)
        elif kind=='code':
            for line in args[0].splitlines():
                p=para(line);p.paragraph_format.space_after=Pt(1)
                for r in p.runs:r.font.name='Courier New';r.font.size=Pt(9)
        elif kind=='picture':pass # Figures are placed deliberately with new cross references.

def build_geometry():
    old.ASSETS=ASSETS
    # Shorten the crowded node label while preserving the original module meaning.
    architecture=(ROOT/'tmp/disclosure/assets/architecture.svg').read_text().replace('page attention proxy','page attention')
    old.save_svg('architecture',[architecture.removesuffix('</svg>')])
    s=old.svg_start(1200,610)
    s.append('<rect x="40" y="95" width="640" height="360" fill="#f5f5f5" stroke="#26394a" stroke-width="2"/>')
    old.text(s,360,65,['Source frame 640 x 360 pixels'],23)
    x,y,w,h=trace[0]['source_crop']
    s.append(f'<rect x="{40+x}" y="{95+y}" width="{w}" height="{h}" fill="#e4eef5" stroke="#156389" stroke-width="3"/>')
    s.append('<rect x="300" y="230" width="85" height="60" fill="#ffe2dc" stroke="#a14231" stroke-width="2"/>')
    old.text(s,343,258,['codec','block'],16);old.arrow(s,'520,275 780,275')
    old.text(s,650,240,['inverse','preprocessing'],17)
    for row in range(14):
        for col in range(14):
            color=['#dfebf5','#e7e1f1','#e3edde'][(row*14+col)//32%3]
            s.append(f'<rect x="{820+col*20}" y="{135+row*20}" width="20" height="20" fill="{color}" stroke="#8f9ca7" stroke-width=".6"/>')
    old.text(s,960,95,['224 x 224 model input'],22)
    old.text(s,960,450,['14 x 14 patch grid','32 patch tokens per page','7 logical pages'],18)
    old.text(s,40,505,['Measured model: Resize 640 x 360 to 440 x 248; crop starts at (108, 12).',f'Source crop: x={x:.2f}, y={y:.2f}, width={w:.2f}, height={h:.2f}.','Logical pages follow token sequence order and can cross image rows.'],19,'start')
    old.save_svg('geometry-measured',s)

def contents():
    question('1 Title of the invention')
    para('System and Method for Generation Safe Motion Guided Visual Key Value Memory Residency and Semantic Correction','title')
    para('Project name: MTR-KV')
    question('2 Field and area of invention')
    para('Computer architecture and systems software; compressed-video processing; visual transformer inference; paged key-value memory; tensor representation management; and versioned runtime control. The invention concerns a processor-based arrangement that translates decoder motion evidence into visual-page residency intent and corrects that intent using later model attention through the same controller.')
    question('3 Prior patents and publications from literature')
    para('The following references identify existing approaches and the technical distinctions proposed for MTR-KV. The focused review was performed on 21 October 2026. It identifies relevant overlap and candidates for examination; it does not establish novelty or freedom to operate. Full source links and publication details appear before Section 10.')
    heading('Prior patent publications')
    # Keep the patent comparison with its original IDF question rather than in an annex.
    emit(2,skip=(1,2,3))
    newpage();heading('Prior research publications')
    emit(1,skip=(0,2))
    heading('Search coverage and remaining work')
    para('The reviewed methods include codec-guided visual-token selection and KV reuse, paging, attention-guided retention, decoder memory and transformer cache management. CodecSight is the closest identified codec-guided reuse lead. The distinctions concern the exact residency transaction, generation checks and model-feedback correction relationship. A claim-by-claim family search, citation expansion and combination analysis remain necessary; an omission from a summary does not establish absence from a reference.')

    newpage();question('4 Summary and background of the invention')
    emit(3,skip=(6,7))
    newpage();heading('4.1 Technical relationship proposed for examination')
    emit(4,skip=(2,3))
    question('5 Objectives of invention')
    emit(5,skip=(6,7,8,9))

    newpage();question('6 Working principle of the invention in brief')
    for kind,args,k in groups[5]['items'][7:]:
        if kind=='p':para(args[0])
    question('7 Description of the invention in detail')
    heading('7.1 System arrangement and component contracts')
    emit(6,skip=(0,1))
    table(['Component','Executable role'],[
        ['MTRE','Motion-to-Token Reconciliation Engine maps source regions to patch tokens and pages.'],
        ['MKAD','Motion-Keyed Allocation Descriptor carries stream, generation, page span and operation.'],
        ['KMC','KV Memory Controller behavioral model owns mappings and readiness.'],
        ['SRCE','Semantic Residency Correction Engine sends model-derived correction through KMC.'],
        ['TorchKVStore','Stores actual fp16, int8 or compressed fp16 K/V associated with page generations.']],[1.2,5.9])
    newpage();picture('architecture','Figure 1 Implemented software architecture',
        'Encoded video produces source-space evidence; MTRE reconciles geometry, temporal state qualifies tier intent, and KMC commits readiness before the model consumes stored K/V. Later attention returns as a same-generation correction through KMC. The dotted GPU/RTL and full VLM extension is proposed.')
    heading('7.2 Execution order and consistency boundary')
    para('The live Python embodiment executes decoder, controller and model work sequentially. Its separate event simulator tests asynchronous ordering and adversarial delays. Logical readiness and a Python fence counter represent the consistency contract; they do not demonstrate a physical GPU fence, asynchronous hardware FIFO or media/GPU overlap.')
    para('A current ready mapping is required before attachment, reuse or read. Semantic correction arrives after that frame has executed; it can change residency safely but cannot avoid the inference pass that produced the score. This timing is part of the implemented control loop.')

    newpage();heading('7.3 Source geometry and page reconciliation')
    picture('geometry-measured','Figure 2 Actual model crop and visual page geometry',
        'The measured ViT transform resizes the 640 by 360 source to 440 by 248, then center-crops 224 by 224 at offset 108 by 12. The inverse source crop is recorded in every trace. Colors illustrate logical pages, which can cross image rows.',
        'full-demo-serial.jsonl source_crop; runtime ViTAttentionFeedback transform Resize 248 and CenterCrop 224.')
    emit(7,skip=(0,3))
    para('The crop in Figure 2 is the benchmark model configuration. A separate regression test deliberately uses Resize 256 to check a different transform; its coordinates are not substituted for the measured run.')

    newpage();heading('7.4 Generation checked allocation and readiness')
    picture('transaction','Figure 3 Transaction and failure preservation',
        'Validate, reserve replacement space, prepare a shadow mapping, and publish before signaling ready. A stale operation is rejected. Failed replacement reservation preserves the active mapping. Atomic publication is modeled in Python.')
    emit(8,skip=(0,))

    newpage();heading('7.5 Actual tensor representations and model feedback')
    emit(10)
    heading('7.6 Exact input validation for temporal reuse')
    para('Low codec motion is an eligibility signal, not a sufficient proof of unchanged model features. Earlier ViT blocks use global attention, so a moving region can change contextual features elsewhere. The final-block adapter compares each eligible page input exactly with the retained feature buffer and additionally checks owner, generation continuity, model signature, geometry, device, dtype and payload age.')
    newpage();picture('reuse','Figure 4 Bounded reuse decision sequence',
        'Readiness failure blocks model access. A failed later reuse condition refreshes the page. Passing candidates rebind retained K/V to the new ready generation; queries and prefix K/V are computed normally, and all patch K/V is still read through the store.')
    emit(9,skip=(0,))

    newpage();heading('7.7 Input transport and optimized implementation')
    emit(11)
    heading('7.8 Other embodiments and industrial application')
    para('The arrangement can be applied to continuous visual inference in edge video analytics, robotics and camera-based assistive systems where visual K/V residency competes with limited fast memory. A hardware embodiment could receive codec metadata from a media engine, translate page intent into device-memory operations and use a hardware completion fence before attention. That embodiment requires its own implementation and validation.')
    para('A full vision-language embodiment would additionally define token ownership across frames and language context, layer-specific K/V layouts and decoder metadata availability. These proposed uses do not extend the evidence from the current final-block image-classifier prototype to all VLM layers or device-level throughput.')

    newpage();question('8 Experimental validation results')
    heading('8.1 Implemented prototype and test setup')
    emit(12,skip=(2,))
    para('Environment: Python 3.14.3, macOS 26.6.2 arm64, PyAV 19.0.1, PyTorch 2.14.1 and timm 1.0.30. The runtime reports CUDA and MPS unavailable. The generated input hash is 5f51727be4c68e07b409f4f56645b84d97a004369ec13fab61d07c359bb36357.')
    heading('8.2 Procedure and comparison')
    para('We ran a 216-frame encoded sequence through the actual decoder, generation-safe controller, tensor store and final ViT attention block. Standard model logits were computed on the same frames for comparison. We also alternated three 60-frame oracle/indexed pairs, excluding the first six frames of each run. The crop, model, storage and reuse settings remained equal; the reconciliation path was the changed condition.')
    para('The timing boundary includes decoder and processing work, including the baseline comparison; it excludes weight loading, JSON serialization, dashboard rendering and display pacing. The primary statistic is the median of three per-run means, not a pooled-frame median.')

    newpage();heading('8.3 Measured optimization results')
    picture('paired-latency','Figure 5 Paired CPU processing time',
        'The indexed path has 52.34 percent lower measured frame time than the reference reconciliation path in the same final build. All paired controller-decision hashes match. This comparison isolates mapping choice and does not establish speedup over standard ViT or a production VLM.',
        'paired-serial/manifest.json and six paired JSONL traces.')
    emit(13,skip=(2,3,4,5,6))

    newpage();heading('8.4 Full sequence validation and interpretation')
    emit(13,skip=(0,1))
    heading('Meaning of agreement and memory figures')
    para('Top-1 agreement means that the stored-K/V classifier and standard classifier selected the same class on these generated frames. It is not task accuracy against ground truth. Tensor payload excludes weights, decoder buffers and Python overhead; the reuse validation buffer is additional memory. There is no measured whole-model memory-reduction claim.')

    newpage();heading('8.5 Behavioral evidence from individual frames')
    para('The following observations come from the same serial trace. They link the claimed mechanism to a trigger, an executed action and a recorded outcome, rather than relying on the interface alone.')
    table(['Frame and trigger','Recorded action','Observed outcome'],[
        ['1\nNo motion-vector metadata','Conservative T0 allocation; no history reused.','READY; 7 recomputed pages; 7 consumed pages; baseline top-1 agrees.'],
        ['3\nQuiet codec evidence','Codec descriptor places pages in T2; semantic feedback promotes pages 0, 5 and 6.','Three APPLIED corrections; final mapping 3 T0 and 4 T2 pages.'],
        ['8\nStatic eligible features','All 7 pages retain origin generation 7 and rebind after current readiness.','196 patch tokens reused; only 1 prefix token projected for K/V.'],
        ['43\nOrigin age exceeded','window_expired for all 7 pages.','0 reused; 7 refreshed. Nearby static frames resume reuse.'],
        ['55\nMoving phase begins','Exact final-block inputs reject 5 candidates; codec/temporal checks reject 2.','All 7 refreshed despite 5 pages passing the earlier eligibility filter.'],
        ['163\nScene-cut phase begins','Missing metadata causes conservative fallback.','All 7 refreshed; no retained payload reused on the cut frame.'],
    ],[1.45,2.75,2.9])
    para('Every listed frame is READY and its baseline top-1 comparison agrees. A selected-record export retains the full descriptors, ready tokens, semantic actions and reuse provenance for inspection. Frame numbers identify processed generations in this fixed file run.')

    newpage();picture('prototype-video-and-pages','Figure 6 Running prototype with committed page residency',
        'At frame 48 the actual dashboard shows the decoded input, seven logical pages, three T0 pages after semantic correction and seven reused K/V pages. The page tiles are logical token-page order, not object boxes. The displayed 33.22 fps is the inverse of this frame\'s decoder/processing time and excludes preview and pacing.',
        'Running local MTR-KV dashboard, 21 October 2026; unmodified prototype-video-and-pages.png. Screenshot session is separate from the serial benchmark.')
    heading('How this view demonstrates the mechanism')
    para('The source pane establishes actual decoding. The committed page display exposes the controller result. Reuse counts come from tensor operations in the model-facing path, while the trace table compares the resulting classifier output with standard inference. The screenshot is corroborating visual evidence; the raw trace remains the primary numerical record.')

    newpage();picture('prototype-reuse-trace','Figure 7 Reuse and age bound visible in the frame table',
        'Frame 43 refreshes all pages at the configured age limit; neighboring static frames reuse seven pages. The displayed observations contain 143186 payload bytes, distinct from the run-wide peak. The quality column reports standard-model top-1 agreement.',
        'Running local MTR-KV dashboard, 21 October 2026; unmodified prototype-reuse-trace.png.')
    heading('8.6 Semantic correction witness')
    f=by_frame[3]
    table(['Page','Codec activity','Attention proxy','Correction','Final tier'],[
        [a['page_id'],f"{a['activity']:.1f}",f"{a['semantic']:.6f}",'T2 to T0','T0'] for a in f['semantic_feedback']['actions']],[.6,1.2,1.3,2,1])
    para('Frame 3 contains a single codec ALLOCATE descriptor covering pages 0 to 6 in T2, followed by three APPLIED semantic corrections for the same generation. The final page records preserve T0 for pages 0, 5 and 6 and T2 for pages 1 to 4. Scores are normalized class-token attention proxies, not human-labeled importance.')

    newpage();heading('8.7 Verification and evidence traceability')
    emit(16)
    para('Selected behavior records are supplied as MTR_KV_Implementation_Records.json with the source trace path and SHA-256. The final document has been checked against those records. The underlying trace, paired comparisons, input and test reports are preserved in the existing evidence archive.')
    heading('Implementation status and operating limits')
    para('The implemented scope is encoded video, actual final-block K/V consumption, generation checks, semantic correction, bounded reuse, a local dashboard and a software-tested Android USB adapter. A physical Android trial and scrcpy/ADB setup remain pending. Generic browser upload, arbitrary-codec success, all-layer VLM processing and real-time performance across devices are not established.')

    newpage();heading('8.8 Reproducing the demonstration')
    emit(14,skip=(0,1,2))
    heading('Evidence locations and entry points')
    emit(21,skip=(2,3,4,5,6))
    para('Source entry points: mtr_kv/mtre.py for reconciliation; controller.py for transactions; live.py for frame and correction coordination; tensor_store.py for actual K/V storage and rebinding; vision_feedback.py for preprocessing and attention integration; video.py for decoding and records. Reproduction requires the dependencies and cached or downloaded model weights.')

    newpage();question('9 Aspects of the invention for which protection is sought')
    emit(17,rename={'Primary coordinated mechanism':'9.1 Coordinated codec intent readiness and correction',
        'Supporting consistency and reuse relationships':'9.2 Failure preservation and qualified temporal reuse',
        'Potential dependent technical details':'9.3 Supporting geometry and generation relationships',
        'What should not be asserted as new alone':'9.4 Technical boundary of the proposal'})
    heading('9.5 Relationship of the implemented evidence to the proposal')
    para('The software establishes a concrete embodiment of the control relationship: frame records show codec tier intent, a current-generation readiness result, actual stored tensor consumption and model-derived correction. Reuse provenance records source and origin generations separately. The examples and tests support narrower technical drafting; they do not determine whether a proposed claim is novel or non-obvious over the cited combinations.')

    newpage();heading('References and reviewed source locations')
    emit(19)
    newpage();heading('References continued and disclosure record')
    emit(20)
    heading('Disclosure and contributors')
    para('The frozen architecture dated 22 September 2026 records Kunwar Shaurya Pratap Singh, Kanan Tayal and Suhani Gupta, B.Tech CSE AI and Data Engineering, VIT Vellore. Contributor names are a design record; legal inventorship, ownership and applicant details require confirmation. The user reported no known earlier public posting, presentation, submission or patent filing on 21 October 2026. Engineering dates and hashes are not patent priority dates.')

    body.insert(len(body)-1,break_proto)
    body.insert(len(body)-1,last_table)
    question('10 Technology readiness level of the invention')
    para('TRL 3 is selected for the demonstrated software mechanism as an experimental proof of concept. The executable prototype processes real encoded video, commits generation-safe page mappings, consumes actual tiered K/V tensors and applies semantic correction. Section 8 supplies the method, measurements, screenshots, frame witnesses and reproducibility records.',bold=True)
    para('This self-assessment applies to the disclosed architecture as a whole. Some software components have local laboratory validation, but physical GPU/RTL integration, full VLM processing and physical-phone operation have not been demonstrated. A user interface alone is not the basis for the selection.')
    para('Evidence date: 21 October 2026. Scope: controlled CPU software proof of concept. Faculty certification is not asserted.')
    body.insert(len(body)-1,trl_table)
    para('Selected column: TRL 3. Inventor contributions, institutional rights, prior-disclosure history and final claim scope should be confirmed with VIT IPR and TT Cell before submission. The present record is an invention disclosure for review; it does not constitute a patent-office filing.')
    para('END OF THE DOCUMENT','sub')

def package():
    OUT.parent.mkdir(parents=True,exist_ok=True)
    buf=io.BytesIO();doc.save(buf);buf.seek(0)
    source=ZipFile(REF);made=ZipFile(buf)
    active=etree.fromstring(made.read('word/document.xml'))
    reltree=etree.fromstring(made.read('word/_rels/document.xml.rels'))
    used={v for el in active.iter() for key,v in el.attrib.items() if key in [qn('r:id'),qn('r:embed'),qn('r:link')]}
    # Keep only active body relationships and all non-image/non-hyperlink structure.
    for rel in list(reltree):
        if rel.get('Type').endswith(('/image','/hyperlink')) and rel.get('Id') not in used:
            reltree.remove(rel)
    keep_media={}
    for name in source.namelist():
        if name.startswith('word/_rels/header') or name.startswith('word/_rels/footer'):
            for rel in etree.fromstring(source.read(name)):
                if rel.get('Type').endswith('/image'):
                    keep_media['word/'+rel.get('Target')]=True
    for rel in reltree:
        if rel.get('Type').endswith('/image'):keep_media['word/'+rel.get('Target')]=True
    changes={'word/document.xml','word/styles.xml','docProps/core.xml','[Content_Types].xml'}
    with ZipFile(OUT,'w',ZIP_DEFLATED) as z:
        for name in source.namelist():
            if name.startswith('word/media/') and name not in keep_media:continue
            value=made.read(name) if name in changes else source.read(name)
            if name=='word/_rels/document.xml.rels':value=etree.tostring(reltree,xml_declaration=True,encoding='UTF-8',standalone=True)
            z.writestr(name,value)
        for name in made.namelist():
            if name not in source.namelist() and (not name.startswith('word/media/') or name in keep_media):z.writestr(name,made.read(name))
    final=ZipFile(OUT)
    inventory=json.loads((TMP/'package-inventory.json').read_text())
    preserved=[]
    for name,info in inventory.items():
        if info['role']=='preserve' and not name.startswith('word/media/'):
            assert sha256(final.read(name)).hexdigest()==info['sha256'],name
            preserved.append(name)
    text=' '.join(etree.fromstring(final.read('word/document.xml')).itertext())
    assert 'AURA' not in text and 'Annex A' not in text
    assert len(Document(OUT).sections)==2
    (TMP/'fidelity-results.json').write_text(json.dumps({'reference_sha256':sha256(REF.read_bytes()).hexdigest(),'preserved_parts':preserved,'output_sha256':sha256(OUT.read_bytes()).hexdigest()},indent=2))
    print(OUT)

if __name__=='__main__':
    build_geometry();contents();package()
